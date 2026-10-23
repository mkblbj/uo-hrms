import importlib
import unittest
from io import BytesIO

from openpyxl import load_workbook

import frappe

from hrms.setup import get_custom_fields


class TestEmployeeWechatField(unittest.TestCase):
	def test_wechat_is_optional_text_for_every_employee(self):
		fields = {field["fieldname"]: field for field in get_custom_fields()["Employee"]}
		self.assertIn("custom_wechat_id", fields)
		field = fields["custom_wechat_id"]
		self.assertEqual(field["fieldtype"], "Data")
		self.assertFalse(field.get("reqd"))
		self.assertFalse(field.get("depends_on"))


class TestEmployeeInformationOverview(unittest.TestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		frappe.db.savepoint("employee_overview_test")
		self.addCleanup(self.cleanup_records)
		self.company = frappe.db.get_value("Company", {"name": "_Test Company"}, "name")
		self.assertTrue(self.company, "Load standard Company test records first")
		self.names = [f"_DIRECTORY-TEST-{i:03d}" for i in range(26)]
		fields = [
			"name",
			"first_name",
			"employee_name",
			"employee_number",
			"company",
			"status",
			"cell_number",
			"current_address",
			"personal_email",
			"company_email",
			"date_of_joining",
			"department",
			"branch",
			"user_id",
			"owner",
		]
		rows = []
		for index, name in enumerate(self.names):
			rows.append(
				[
					name,
					f"Directory {index:02d}",
					f"Directory {index:02d}",
					f"{index:06d}",
					self.company,
					["Active", "Left", "Inactive", "Suspended"][index % 4],
					"09001234567" if index == 0 else "+81-90-1111-2222",
					"東京都千代田区\n建物 101",
					f"directory{index}@example.com",
					"",
					"2026-01-01",
					"_Directory North",
					"_Directory Tokyo",
					"",
					"Administrator",
				]
			)
		frappe.db.bulk_insert("Employee", fields, rows)
		# Exercise NULL, empty strings and whitespace independently.
		for index, phone, address in [(1, None, None), (2, "", ""), (3, " \t ", " \n ")]:
			frappe.db.set_value(
				"Employee", self.names[index], {"cell_number": phone, "current_address": address}
			)
		if frappe.get_meta("Employee").has_field("custom_wechat_id"):
			frappe.db.set_value("Employee", self.names[0], "custom_wechat_id", "wx_00123")

	def cleanup_records(self):
		frappe.set_user("Administrator")
		frappe.db.rollback(save_point="employee_overview_test")
		frappe.clear_cache()

	def report(self, **filters):
		module = importlib.import_module(
			"hrms.hr.report.employee_information_overview.employee_information_overview"
		)
		return module.execute({"search": "_DIRECTORY-TEST-", "department": "_Directory North", **filters})

	def make_user(self, role):
		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": f"directory-{role.lower().replace(' ', '-')}@example.com",
				"first_name": "Directory Test",
				"send_welcome_email": 0,
				"roles": [{"role": role}],
			}
		).insert(ignore_permissions=True)
		return user.name

	def test_all_statuses_and_employees_without_accounts_are_returned_without_page_limit(self):
		columns, data = self.report()
		self.assertEqual(len(data), 26)
		self.assertEqual({row["name"] for row in data}, set(self.names))
		self.assertEqual({row["status"] for row in data}, {"Active", "Left", "Inactive", "Suspended"})
		self.assertEqual([row["employee_name"] for row in data], sorted(row["employee_name"] for row in data))
		column_names = {column["fieldname"] for column in columns}
		self.assertTrue(
			{"name", "employee_name", "cell_number", "custom_wechat_id", "current_address"} <= column_names
		)

	def test_search_matches_name_number_phone_and_wechat(self):
		for search in ["Directory 00", "000000", "09001234567", "wx_00123"]:
			with self.subTest(search=search):
				self.assertEqual([row["name"] for row in self.report(search=search)[1]], self.names[:1])

	def test_search_treats_percent_underscore_and_quotes_as_literal_text(self):
		frappe.db.set_value("Employee", self.names[0], "custom_wechat_id", "wx_%'literal")
		self.assertEqual(len(self.report(search="wx_%'literal")[1]), 1)
		self.assertEqual(self.report(search="%' OR 1=1 --")[1], [])

	def test_combined_filters_and_missing_phone_or_address(self):
		self.assertEqual(
			len(self.report(status="Active", department="_Directory North", company=self.company)[1]), 7
		)
		for flag in ("missing_phone", "missing_address"):
			self.assertEqual({row["name"] for row in self.report(**{flag: 1})[1]}, set(self.names[1:4]))
		self.assertEqual(len(self.report(missing_phone="0", status="Active")[1]), 7)
		self.assertEqual(
			[row["name"] for row in self.report(missing_address=1, status="Left")[1]], self.names[1:2]
		)

	def test_email_falls_back_to_linked_account_and_preserves_employee_email_priority(self):
		self.make_user("Employee")
		user = "directory-employee@example.com"
		frappe.db.set_value(
			"Employee", self.names[0], {"user_id": user, "company_email": " \t ", "personal_email": ""}
		)
		self.assertEqual(self.report(search=self.names[0])[1][0]["contact_email"], user)
		frappe.db.set_value("Employee", self.names[0], "personal_email", "personal@example.com")
		self.assertEqual(self.report(search=self.names[0])[1][0]["contact_email"], "personal@example.com")
		frappe.db.set_value("Employee", self.names[0], "company_email", "company@example.com")
		self.assertEqual(self.report(search=self.names[0])[1][0]["contact_email"], "company@example.com")
		frappe.db.set_value(
			"Employee", self.names[0], {"user_id": "Administrator", "company_email": "", "personal_email": ""}
		)
		self.assertEqual(self.report(search=self.names[0])[1][0]["contact_email"], "")

	def test_branch_is_not_returned_or_exported_and_email_is_easy_to_find(self):
		columns, data = self.report()
		fields = [column["fieldname"] for column in columns]
		self.assertNotIn("branch", fields)
		self.assertTrue(all("branch" not in row for row in data))
		self.assertLess(fields.index("contact_email"), fields.index("cell_number"))

	def test_contact_values_stay_raw_and_updates_use_the_same_employee(self):
		columns, data = self.report(search=self.names[0])
		self.assertEqual(data[0]["cell_number"], "09001234567")
		self.assertEqual(data[0]["custom_wechat_id"], "wx_00123")
		self.assertEqual(data[0]["current_address"], "東京都千代田区\n建物 101")
		self.assertEqual(data[0]["contact_email"], "directory0@example.com")
		frappe.db.set_value("Employee", self.names[0], "current_address", "大阪府\n新住所")
		self.assertEqual(self.report(search=self.names[0])[1][0]["current_address"], "大阪府\n新住所")
		self.assertEqual(len(self.report()[1]), 26)
		for column in columns:
			if column["fieldname"] in {"cell_number", "custom_wechat_id", "employee_number"}:
				self.assertEqual(column["fieldtype"], "Data")

	def test_guest_and_employee_are_denied_before_employee_query(self):
		for user in ["Guest", self.make_user("Employee")]:
			with self.subTest(user=user):
				frappe.set_user(user)
				with self.assertRaises(frappe.PermissionError):
					self.report()
			frappe.set_user("Administrator")

	def test_hr_roles_can_read_and_existing_user_permissions_are_preserved(self):
		for role in ["HR User", "HR Manager"]:
			user = self.make_user(role)
			frappe.set_user(user)
			self.assertEqual(len(self.report()[1]), 26, msg=role)
			frappe.set_user("Administrator")
		user = "directory-hr-user@example.com"
		frappe.get_doc(
			{
				"doctype": "User Permission",
				"user": user,
				"allow": "Employee",
				"for_value": self.names[0],
				"apply_to_all_doctypes": 1,
			}
		).insert(ignore_permissions=True)
		frappe.clear_cache(user=user)
		frappe.set_user(user)
		self.assertEqual([row["name"] for row in self.report()[1]], self.names[:1])

	def test_system_manager_must_also_have_existing_employee_permissions(self):
		user = self.make_user("System Manager")
		frappe.set_user(user)
		with self.assertRaises(frappe.PermissionError):
			self.report()
		frappe.set_user("Administrator")
		frappe.get_doc("User", user).add_roles("HR User")
		frappe.set_user(user)
		self.assertEqual(len(self.report()[1]), 26)

	def test_report_export_contains_all_filtered_rows_and_text_values(self):
		from frappe.desk.query_report import _export_query
		from frappe.desk.query_report import run as run_report

		emergency = {
			"person_to_be_contacted": "李四",
			"relation": "配偶",
			"emergency_phone_number": "09000001234",
		}
		frappe.db.set_value("Employee", self.names[0], dict(emergency))

		result = run_report("Employee Information Overview", {"search": "_DIRECTORY-TEST-"})
		self.assertEqual(len(result["result"]), 26)
		for field, value in emergency.items():
			self.assertEqual(result["result"][0][field], value)
		_params = frappe._dict(
			{
				"report_name": "Employee Information Overview",
				"file_format_type": "Excel",
				"filters": {"search": "_DIRECTORY-TEST-"},
			}
		)
		_, extension, content = _export_query(_params, {}, populate_response=False)
		self.assertEqual(extension, "xlsx")
		sheet = load_workbook(BytesIO(content)).active
		self.assertEqual(sheet.max_row, 27)
		self.assertIn("09001234567", [cell.value for cell in sheet[2]])
		self.assertIn("000000", [cell.value for cell in sheet[2]])
		self.assertIn("wx_00123", [cell.value for cell in sheet[2]])
		self.assertIn("東京都千代田区\n建物 101", [cell.value for cell in sheet[2]])
		for value in emergency.values():
			self.assertIn(value, [cell.value for cell in sheet[2]])

	def test_employee_cannot_run_or_export_report(self):
		from frappe.desk.query_report import export_query
		from frappe.desk.query_report import run as run_report

		frappe.set_user(self.make_user("Employee"))
		with self.assertRaises(frappe.PermissionError):
			run_report("Employee Information Overview", {})
		frappe.local.form_dict = frappe._dict(
			{"report_name": "Employee Information Overview", "file_format_type": "Excel"}
		)
		with self.assertRaises(frappe.PermissionError):
			export_query()

	def editable_employee(self):
		from frappe.client import get

		frappe.db.set_value(
			"Employee",
			self.names[0],
			{
				"gender": "Female",
				"date_of_birth": "1990-01-01",
				"department": "",
				"branch": "",
				"naming_series": frappe.get_meta("Employee")
				.get_field("naming_series")
				.options.splitlines()[0],
			},
		)
		return get("Employee", self.names[0])

	def test_quick_edit_native_save_updates_contacts_without_creating_another_employee(self):
		from frappe.client import save

		doc = self.editable_employee()
		frappe.set_user(self.make_user("HR User"))
		values = {
			"cell_number": "+81-90-0123-4567",
			"custom_wechat_id": "quick_edit_wx",
			"current_address": "大阪府\n新住所",
			"company_email": "work@example.com",
			"personal_email": "personal@example.com",
			"person_to_be_contacted": "李四",
			"relation": "配偶",
			"emergency_phone_number": "09000001234",
		}
		doc.update(values)
		save(doc)
		stored = frappe.get_doc("Employee", self.names[0])
		for field, value in values.items():
			self.assertEqual(stored.get(field), value)
		self.assertEqual(stored.status, "Active")
		self.assertEqual(frappe.db.count("Employee", {"name": ["like", "_DIRECTORY-TEST-%"]}), 26)

	def test_quick_edit_native_save_requires_write_permission(self):
		from frappe.client import save

		doc = self.editable_employee()
		frappe.set_user(self.make_user("Employee"))
		doc["cell_number"] = "09000000000"
		with self.assertRaises(frappe.PermissionError):
			save(doc)
		frappe.set_user("Administrator")
		self.assertEqual(frappe.db.get_value("Employee", self.names[0], "cell_number"), "09001234567")

	def test_quick_edit_native_save_rejects_a_stale_employee(self):
		from frappe.client import save

		doc = self.editable_employee()
		frappe.db.set_value("Employee", self.names[0], "current_address", "由其他人更新的住址")
		doc["current_address"] = "旧窗口的新输入"
		with self.assertRaises(frappe.TimestampMismatchError):
			save(doc)
		self.assertEqual(
			frappe.db.get_value("Employee", self.names[0], "current_address"), "由其他人更新的住址"
		)
